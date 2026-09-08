from __future__ import annotations
import json, subprocess
from pathlib import Path
import pytest
from pptx import Presentation
from PIL import Image
from test_runtime_generation import ROOT, PY, invoke

@pytest.mark.parametrize("mode",["quiz","flashcard","lottery","timer"])
def test_mini_app_modes_and_injection(mode,tmp_path):
    skill="tw-edu-mini-app"; d=json.loads((ROOT/skill/"examples/example.json").read_text()); c={"title":"<img src=x onerror=alert(1)>","mode":mode}
    if mode=="quiz": c["questions"]=[{"id":"q","prompt":"<script>x</script>","points":2,"options":[{"id":"a","text":"A"},{"id":"b","text":"B"}],"answer":"a"}]
    if mode=="flashcard": c["cards"]=[{"front":"<b>前</b>","back":"後"}]
    if mode=="lottery": c["entries"]=["甲","乙"]
    if mode=="timer": c["seconds"]=3
    d["content"]=c; inp=tmp_path/"in.json"; inp.write_text(json.dumps(d)); out=tmp_path/"x.html"; r=invoke(skill,"--input",inp,"--output",out); assert r.returncode==0,r.stderr
    page=out.read_text(); assert "innerHTML" not in page and "\\u003c" in page and "<script>x</script>" not in page
    if mode=="quiz": assert "得分：" in page and "重新作答" in page and "解析：" in page

def test_native_slides_have_text_table_chart_notes(tmp_path):
    out=tmp_path/"x.pptx"; r=invoke("tw-edu-slides-creator","--example","--output",out); assert r.returncode==0,r.stderr
    prs=Presentation(out); assert len(prs.slides)==3; assert any(s.has_table for s in prs.slides[1].shapes); assert any(s.has_chart for s in prs.slides[2].shapes); assert "slide_id:1" in prs.slides[0].notes_slide.notes_text_frame.text
    assert prs.slide_width/prs.slide_height==pytest.approx(16/9,rel=.001)

def test_image_slide_gap_and_aspect(tmp_path):
    script=ROOT/"tw-edu-slides-creator/scripts/assemble_image_pptx.py"; deck=tmp_path/"deck"; origin=deck/"origin_image"; origin.mkdir(parents=True)
    Image.new("RGB",(1600,900)).save(origin/"slide_01.png"); Image.new("RGB",(1600,900)).save(origin/"slide_03.png")
    r=subprocess.run([str(PY),str(script),str(tmp_path),"deck.pptx"],text=True,capture_output=True); assert r.returncode==2 and "contiguous" in r.stderr

def test_validate_only_no_output(tmp_path):
    skill="tw-edu-worksheet-creator"; out=tmp_path/"x.docx"; r=invoke(skill,"--example","--validate-only"); assert r.returncode==0 and not out.exists()
