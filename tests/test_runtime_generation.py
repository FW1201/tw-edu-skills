from __future__ import annotations
import json, os, subprocess, sys
from pathlib import Path
import pytest

ROOT=Path(__file__).resolve().parents[1]
PY=Path(sys.executable)
MANIFEST=json.loads((ROOT/"skills-manifest.json").read_text(encoding="utf-8"))
GENERATORS={s["name"]:Path(s["entrypoint"]).name for s in MANIFEST["skills"] if s["entrypoint"]}
EXT={s["name"]:"."+s["format"] for s in MANIFEST["skills"] if s["entrypoint"] and s["format"]!="docx"}
def invoke(skill,*args,cwd=None):
    env=os.environ.copy(); env.setdefault("MPLCONFIGDIR","/private/tmp/tw-edu-mpl"); env.setdefault("XDG_CACHE_HOME","/private/tmp/tw-edu-cache")
    return subprocess.run([str(PY),str(ROOT/skill/"scripts"/GENERATORS[skill]),*map(str,args)],cwd=cwd,text=True,capture_output=True,env=env)

def test_manifest_lists_every_generator():
    assert len(GENERATORS)==33 and "tw-edu-official-document" in GENERATORS

@pytest.mark.parametrize("skill",GENERATORS)
def test_every_generator_example_and_arbitrary_cwd(skill,tmp_path):
    out=tmp_path/("result"+EXT.get(skill,".docx")); r=invoke(skill,"--example","--output",out,cwd=tmp_path)
    assert r.returncode==0,r.stderr
    if skill=="tw-edu-exam-generator":
        assert (tmp_path/"result-student.docx").exists() and (tmp_path/"result-teacher.docx").exists()
    else: assert out.exists()

@pytest.mark.parametrize("skill",GENERATORS)
def test_bad_input_creates_no_output(skill,tmp_path):
    bad=tmp_path/"bad.json"; bad.write_text('{}'); out=tmp_path/("bad"+EXT.get(skill,".docx"))
    r=invoke(skill,"--input",bad,"--output",out); assert r.returncode!=0; assert not out.exists()

def test_exam_student_has_no_answers(tmp_path):
    from docx import Document
    out=tmp_path/"exam.docx"; r=invoke("tw-edu-exam-generator","--example","--output",out); assert r.returncode==0
    student="\n".join(p.text for p in Document(tmp_path/"exam-student.docx").paragraphs)
    teacher="\n".join(p.text for p in Document(tmp_path/"exam-teacher.docx").paragraphs)
    assert "答案：A" not in student and "答案：A" in teacher

def test_lesson_bad_minutes_no_output(tmp_path):
    skill="tw-edu-lesson-plan-108"; data=json.loads((ROOT/skill/"examples/example.json").read_text()); data["content"]["total_minutes"]=99
    inp=tmp_path/"x.json"; inp.write_text(json.dumps(data)); out=tmp_path/"x.docx"; r=invoke(skill,"--input",inp,"--output",out)
    assert r.returncode!=0 and not out.exists()

def test_input_content_changes_output(tmp_path):
    from docx import Document
    skill="tw-edu-feedback-writer"; data=json.loads((ROOT/skill/"examples/example.json").read_text()); data["context"]={"subject":"數學","grade":"高中二年級","topic":"二次函數"}; data["content"]["students"][0]["observations"]=["學生能由圖形判讀頂點"]
    inp=tmp_path/"x.json"; inp.write_text(json.dumps(data,ensure_ascii=False)); out=tmp_path/"x.docx"; assert invoke(skill,"--input",inp,"--output",out).returncode==0
    doc=Document(out); text="\n".join([*(p.text for p in doc.paragraphs),*(cell.text for table in doc.tables for row in table.rows for cell in row.cells)])
    assert "學生能由圖形判讀頂點" in text and "高中二年級" in text

def test_task_specific_schemas_are_not_generic_blocks():
    expected={"tw-edu-feedback-writer":"students","tw-edu-learning-portfolio":"records","tw-edu-parent-communication":"recipients","tw-edu-formative-assessment":"checks","tw-edu-worksheet-creator":"prompts"}
    for skill,field in expected.items():
        schema=json.loads((ROOT/skill/"schemas/input.schema.json").read_text()); required=schema["properties"]["content"]["required"]
        assert field in required and "sections" not in required
