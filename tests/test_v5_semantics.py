"""v5 semantic gates: every new or restructured skill accepts its example and blocks known-bad inputs."""
from __future__ import annotations
import copy, json, subprocess, sys
from pathlib import Path
import pytest

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = json.loads((ROOT / 'skills-manifest.json').read_text(encoding='utf-8'))
ENTRY = {s['name']: s['entrypoint'] for s in MANIFEST['skills']}


def example(skill):
    return json.loads((ROOT / skill / 'examples/example.json').read_text(encoding='utf-8'))


def validate(skill, data, tmp_path):
    path = tmp_path / 'input.json'; path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    return subprocess.run([sys.executable, str(ROOT / skill / ENTRY[skill]), '--input', str(path), '--validate-only'], text=True, capture_output=True)


def generate(skill, data, tmp_path):
    tmp_path.mkdir(parents=True, exist_ok=True)
    path = tmp_path / 'input.json'; path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    out = tmp_path / 'out.docx'
    r = subprocess.run([sys.executable, str(ROOT / skill / ENTRY[skill]), '--input', str(path), '--output', str(out)], text=True, capture_output=True)
    assert r.returncode == 0, r.stderr
    return out, json.loads(out.with_name(out.name + '.validation.json').read_text(encoding='utf-8'))


def docx_text(path):
    from docx import Document
    doc = Document(path)
    return '\n'.join([*(p.text for p in doc.paragraphs), *(c.text for t in doc.tables for r in t.rows for c in r.cells)])


V5 = ['tw-edu-homeroom-operations', 'tw-edu-behavior-support', 'tw-edu-incident-response', 'tw-edu-student-guidance-advice',
      'tw-edu-guidance-collaboration', 'tw-edu-official-document', 'tw-edu-lesson-design-brainstorm', 'tw-edu-school-curriculum-plan',
      'tw-edu-open-lesson', 'tw-edu-conduct-comments', 'tw-edu-school-affairs-meeting', 'tw-edu-teacher-wellbeing',
      'tw-edu-classroom-culture', 'tw-edu-parent-communication', 'tw-edu-school-document', 'tw-edu-meeting-facilitator']


@pytest.mark.parametrize('skill', V5)
def test_example_passes(skill, tmp_path):
    r = validate(skill, example(skill), tmp_path); assert r.returncode == 0, r.stderr


def mutate(skill, fn):
    d = copy.deepcopy(example(skill)); fn(d['content']); return d


NEGATIVE = [
    ('tw-edu-homeroom-operations', lambda c: c.update(period='學期末交接'), 'handover'),
    ('tw-edu-homeroom-operations', lambda c: c['roles'].append(dict(c['roles'][0])), 'duplicate role'),
    ('tw-edu-behavior-support', lambda c: c.update(student_alias='王小明'), 'anonymous'),
    ('tw-edu-behavior-support', lambda c: c['response_plan'].append({'situation': '再犯', 'teacher_response': '罰站一節課'}), 'punishment'),
    ('tw-edu-behavior-support', lambda c: c['safety_check'].update(red_flags_present=True), 'red flags'),
    ('tw-edu-behavior-support', lambda c: c['function_hypotheses'][0].update(statement='學生有過動症所以離座'), 'diagnose'),
    ('tw-edu-incident-response', lambda c: c.update(case_status='已結案'), 'cannot be closed'),
    ('tw-edu-incident-response', lambda c: c.pop('reporting'), 'reporting route'),
    ('tw-edu-incident-response', lambda c: c['restorative'].update(eligible=True), 'formal procedure'),
    ('tw-edu-student-guidance-advice', lambda c: c['analysis']['ai_inferences'].append('可能是憂鬱症'), 'diagnose'),
    ('tw-edu-student-guidance-advice', lambda c: c['red_flag_screen'].update(triggered=True, flags=['自傷言語']) or c['referral'].update(recommended=False), 'referral'),
    ('tw-edu-student-guidance-advice', lambda c: c.update(hypotheses=c['hypotheses'][:1]), 'validation failed'),
    ('tw-edu-guidance-collaboration', lambda c: c.pop('meeting_agenda'), 'IEP'),
    ('tw-edu-guidance-collaboration', lambda c: c.update(concern='疑似自閉症'), 'diagnose'),
    ('tw-edu-official-document', lambda c: c.update(subject='一、檢陳成果報告\n二、請鑒核'), '主旨'),
    ('tw-edu-official-document', lambda c: c.update(direction='平行'), '鈞'),
    ('tw-edu-official-document', lambda c: c.update(subject=c['subject'].replace('請鑒核', '請 鑒核')), '挪抬'),
    ('tw-edu-official-document', lambda c: c.update(doc_type='簽', direction='校內'), '擬辦'),
    ('tw-edu-lesson-design-brainstorm', lambda c: c['learning_focus'][0].pop('verification_note'), 'verification_note'),
    ('tw-edu-lesson-design-brainstorm', lambda c: c.update(next_skill='tw-edu-unknown'), 'validation failed'),
    ('tw-edu-school-curriculum-plan', lambda c: c['required_periods'][0].update(periods_per_week=5), 'differ'),
    ('tw-edu-school-curriculum-plan', lambda c: c['goals'].append({'id': 'G9', 'goal': '未被承擔的目標', 'image_links': ['能合作的公民']}), 'unused'),
    ('tw-edu-school-curriculum-plan', lambda c: c['units'][0].update(course='不存在的課'), 'not in structure'),
    ('tw-edu-open-lesson', lambda c: c['observations'][0].update(behavior='第1組討論熱烈，表現很好'), 'low-inference'),
    ('tw-edu-open-lesson', lambda c: c['prep'].update(observation_focus=['一', '二', '三', '四']), 'validation failed'),
    ('tw-edu-open-lesson', lambda c: c.pop('debrief'), 'debrief'),
    ('tw-edu-conduct-comments', lambda c: c['students'][0].update(evidence_ids=['E9']), 'unknown evidence'),
    ('tw-edu-conduct-comments', lambda c: c['students'][0].update(comment='表現優良，評為甲等'), '等第'),
    ('tw-edu-conduct-comments', lambda c: c['students'][0].update(comment='上課懶惰，需要加強'), 'labelling'),
    ('tw-edu-school-affairs-meeting', lambda c: c['allegations'].append({'id': 'A2', 'summary': '另一事項', 'my_account': '說明', 'evidence_available': [], 'evidence_to_request': []}), 'every allegation'),
    ('tw-edu-school-affairs-meeting', lambda c: c.update(support_people=[{'who': f'人員{i}', 'role': '輔佐人'} for i in range(3)]), '輔佐人'),
    ('tw-edu-school-affairs-meeting', lambda c: c.pop('interview_prep'), 'interview_prep'),
    ('tw-edu-teacher-wellbeing', lambda c: c['risk_screen'].update(self_harm_thoughts=True) or c.update(resources=[{'resource': '同事', 'access': '聊天'}]), 'crisis'),
    ('tw-edu-teacher-wellbeing', lambda c: c['risk_screen'].update(unable_to_function=True) or c.update(plan=[{'horizon': '本月', 'action': '休假'}]), 'action now'),
    ('tw-edu-classroom-culture', lambda c: c['response_plan'].append({'trigger': '忘記帶作業', 'teacher_action': '全班一起罰抄', 'follow_up': '無'}), 'punishment'),
    ('tw-edu-classroom-culture', lambda c: c['repair_paths'][0]['steps'].append('罰錢新臺幣50元'), 'punishment'),
    ('tw-edu-parent-communication', lambda c: c.update(audience='全班家長'), 'all parents'),
    ('tw-edu-parent-communication', lambda c: c.update(sensitivity='正式程序', channel='私訊'), '正式程序'),
    ('tw-edu-school-document', lambda c: c.update(total_budget=99999), 'sum'),
    ('tw-edu-school-document', lambda c: c.update(document_type='成果報告'), 'results'),
    ('tw-edu-meeting-facilitator', lambda c: c.update(duration_minutes=90), 'duration'),
    ('tw-edu-meeting-facilitator', lambda c: c['discussions'][0].update(resolution='照案通過'), '待決議'),
]


@pytest.mark.parametrize('skill,change,message', NEGATIVE, ids=[f'{s}-{i}' for i, (s, _, _) in enumerate(NEGATIVE)])
def test_semantic_gate_blocks(skill, change, message, tmp_path):
    r = validate(skill, mutate(skill, change), tmp_path)
    assert r.returncode != 0, f'{skill} accepted bad input'
    assert message in r.stderr, r.stderr


def test_every_v5_skill_has_two_negatives():
    counts = {s: sum(1 for x, _, _ in NEGATIVE if x == s) for s in V5}
    assert all(n >= 2 for n in counts.values()), counts


def test_negated_phrases_do_not_trigger(tmp_path):
    d = mutate('tw-edu-classroom-culture', lambda c: c['norms']['safety_floor'].append('教師不得體罰，也不當眾羞辱學生'))
    assert validate('tw-edu-classroom-culture', d, tmp_path).returncode == 0


def test_official_document_layout_and_advisories(tmp_path):
    out, report = generate('tw-edu-official-document', example('tw-edu-official-document'), tmp_path)
    text = docx_text(out)
    for part in ['○○市立○○國民中學　函', '受文者：○○市政府教育局', '主旨：檢陳', '說明：', '一、依鈞局', '正本：', '附註（不屬公文本文）']:
        assert part in text, part
    assert any('草稿' in a for a in report['advisories'])
    d = mutate('tw-edu-official-document', lambda c: c.update(proposal_label='辦法', proposal=['惠請貴局核撥經費']))
    _, report = generate('tw-edu-official-document', d, tmp_path / 'b')
    assert any('第三段' in a for a in report['advisories']) and any('惠請' in a for a in report['advisories']), report


def test_incident_record_keeps_reporting_and_not_alone(tmp_path):
    out, report = generate('tw-edu-incident-response', example('tw-edu-incident-response'), tmp_path)
    text = docx_text(out)
    assert '校園霸凌防制準則第17條' in text and '不安排兩人和解或當面對質' in text and '已通報待交接' in text
    assert any('待通報' in a for a in report['advisories'])


def test_input_content_reaches_v5_documents(tmp_path):
    d = mutate('tw-edu-student-guidance-advice', lambda c: c.update(teacher_view='我擔心他週末打工太累'))
    out, _ = generate('tw-edu-student-guidance-advice', d, tmp_path)
    assert '我擔心他週末打工太累' in docx_text(out)
    d = mutate('tw-edu-conduct-comments', lambda c: c['students'][0].update(suggestion='下學期試著擔任小組長'))
    out, _ = generate('tw-edu-conduct-comments', d, tmp_path / 'c')
    text = docx_text(out)
    assert '下學期試著擔任小組長' in text and 'E1' in text


def test_curriculum_mapper_events_sheet(tmp_path):
    from openpyxl import load_workbook
    d = example('tw-edu-curriculum-mapper'); d['content']['school_events'] = [{'week': '第7週', 'event': '第一次段考'}]
    path = tmp_path / 'input.json'; path.write_text(json.dumps(d, ensure_ascii=False), encoding='utf-8')
    out = tmp_path / 'map.xlsx'
    r = subprocess.run([sys.executable, str(ROOT / 'tw-edu-curriculum-mapper' / ENTRY['tw-edu-curriculum-mapper']), '--input', str(path), '--output', str(out)], text=True, capture_output=True)
    assert r.returncode == 0, r.stderr
    wb = load_workbook(out); assert wb['校行事對照']['B2'].value == '第一次段考'
