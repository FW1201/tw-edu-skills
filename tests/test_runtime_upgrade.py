import json
from pathlib import Path
import pytest
from PIL import Image
from pptx import Presentation
from test_runtime_generation import ROOT, invoke


def payload(skill):
    return json.loads((ROOT / skill / 'examples/example.json').read_text())


def render(skill, data, tmp_path, suffix='.docx', validate=False):
    source = tmp_path / 'input.json'
    source.write_text(json.dumps(data, ensure_ascii=False))
    output = tmp_path / ('result' + suffix)
    args = ['--input', source, '--validate-only'] if validate else ['--input', source, '--output', output]
    return invoke(skill, *args, cwd=tmp_path), output


def test_image_mode_and_notes(tmp_path):
    skill = 'tw-edu-slides-creator'
    data = payload(skill)
    Image.new('RGB', (1600, 900), 'navy').save(tmp_path / 'page.png')
    data['content'] = {'title': '圖像簡報', 'mode': 'image', 'slides': [{'id': 1, 'type': 'image', 'title': '圖片頁', 'image': 'page.png', 'notes': '同頁講者備註'}]}
    result, output = render(skill, data, tmp_path, '.pptx')
    assert result.returncode == 0, result.stderr
    prs = Presentation(output)
    assert len(prs.slides[0].shapes) == 1
    assert '同頁講者備註' in prs.slides[0].notes_slide.notes_text_frame.text


@pytest.mark.parametrize('defect', ['missing', 'aspect', 'duplicate', 'gap', 'ragged'])
def test_slide_preflight_rejects_before_output(tmp_path, defect):
    skill = 'tw-edu-slides-creator'
    data = payload(skill)
    if defect == 'ragged':
        data['content']['slides'][1]['table'][1].append('extra')
    else:
        Image.new('RGB', (100, 100) if defect == 'aspect' else (1600, 900)).save(tmp_path / 'page.png')
        slide = {'id': 1, 'type': 'image', 'title': '圖', 'image': 'missing.png' if defect == 'missing' else 'page.png'}
        data['content'] = {'title': '圖片', 'mode': 'image', 'slides': [slide]}
        if defect == 'duplicate': data['content']['slides'].append(dict(slide, id=2))
        if defect == 'gap': slide['id'] = 2
    result, output = render(skill, data, tmp_path, '.pptx')
    assert result.returncode == 2 and not output.exists()
    assert 'Traceback' not in result.stderr


@pytest.mark.parametrize('field', ['dimensions', 'levels'])
def test_invalid_analytic_rubric(tmp_path, field):
    skill = 'tw-edu-rubric-designer'
    data = payload(skill)
    if field == 'dimensions': del data['content'][field]
    else: data['content'][field][1]['id'] = data['content'][field][0]['id']
    result, output = render(skill, data, tmp_path)
    assert result.returncode == 2 and not output.exists()
    assert 'Traceback' not in result.stderr


def test_holistic_rubric(tmp_path):
    from docx import Document
    skill = 'tw-edu-rubric-designer'
    data = payload(skill)
    c = data['content']; c['type'] = 'holistic'; del c['dimensions']
    c['total_points'] = 2
    c['descriptions'] = {'L1': '整體论證有證據', 'L0': '整體論證待補強'}
    result, output = render(skill, data, tmp_path)
    assert result.returncode == 0, result.stderr
    assert len(Document(output).tables[0].columns) == 3


@pytest.mark.parametrize('value', ['   ', '\n'])
def test_whitespace_content_rejected(tmp_path, value):
    skill = 'tw-edu-lesson-plan-108'
    data = payload(skill); data['content']['title'] = value
    result, output = render(skill, data, tmp_path)
    assert result.returncode == 2 and not output.exists()


def test_duplicate_quiz_options(tmp_path):
    skill = 'tw-edu-mini-app'
    data = payload(skill)
    data['content']['questions'][0]['options'][1]['id'] = 'A'
    result, output = render(skill, data, tmp_path, '.html')
    assert result.returncode == 2 and not output.exists()


def test_existing_output_is_preserved(tmp_path):
    skill = 'tw-edu-lesson-plan-108'
    out = tmp_path / 'result.docx'; out.write_bytes(b'custom')
    result, _ = render(skill, payload(skill), tmp_path)
    assert result.returncode == 2 and out.read_bytes() == b'custom'
