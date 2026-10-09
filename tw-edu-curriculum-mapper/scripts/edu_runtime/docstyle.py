"""Shared DOCX page setup and finishing styles."""
from __future__ import annotations

from pathlib import Path

FONT = "Noto Sans TC"
_WEIGHTS: dict = {}


def new_document():
    from docx import Document
    from docx.shared import Cm
    doc = Document(); sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8); sec.left_margin = sec.right_margin = Cm(2)
    return doc


def add_table(doc, headers, rows, weights=None):
    table = doc.add_table(rows=1, cols=len(headers)); table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, headers): cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row): cell.text = text
    if weights: _WEIGHTS[id(table._tbl)] = (table._tbl, weights)
    return table


def finish(doc, output: Path) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    styles = doc.styles
    for n in ["Normal", "Title", "Heading 1", "Heading 2"]:
        styles[n].font.name = FONT; styles[n]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT); styles[n].font.size = Pt(11 if n == "Normal" else 16)
        styles[n].paragraph_format.line_spacing = 1.35
        styles[n].font.color.rgb = RGBColor(0, 0, 0)
        for border in styles[n]._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for paragraph in doc.paragraphs:
        for border in paragraph._p.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    for table in doc.tables:
        header = OxmlElement('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(header)
        headings = [cell.text for cell in table.rows[0].cells]
        weights = _WEIGHTS.get(id(table._tbl), (None, None))[1] or [1 if h in {'編號', '分鐘', '配分'} else 2 if h == '對應目標' else 4 for h in headings]
        table.autofit = False
        for j, weight in enumerate(weights):
            table.columns[j].width = Cm(17 * weight / sum(weights))
        for row in table.rows:
            for j, cell in enumerate(row.cells):
                cell.width = Cm(17 * weights[j] / sum(weights))
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        run.font.name = FONT
                        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = FONT; run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)
