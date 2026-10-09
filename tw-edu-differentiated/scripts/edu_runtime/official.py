"""Renderers with fixed document conventions: official documents and conduct comments."""
from __future__ import annotations

from .docstyle import add_table

DIGITS = '〇一二三四五六七八九'


def zh_num(n: int) -> str:
    if n < 10: return DIGITS[n]
    tens, ones = divmod(n, 10)
    return ('' if tens == 1 else DIGITS[tens]) + '十' + (DIGITS[ones] if ones else '')


def _numbered(doc, items):
    if len(items) == 1:
        doc.add_paragraph(items[0]); return
    for i, text in enumerate(items, 1): doc.add_paragraph(f'{zh_num(i)}、{text}')


def _line(doc, label, value):
    if value: doc.add_paragraph(f'{label}：{value}')


def render_official(doc, d):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    c = d['content']; kind = c['doc_type']
    heading = {'簽': f"簽　於{c['issuer']}", '便簽': f"{c['issuer']}　便簽", '公告': f"{c['issuer']}　公告", '開會通知單': f"{c['issuer']}　開會通知單"}.get(kind, f"{c['issuer']}　{kind}")
    p = doc.add_heading(heading, 0); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if kind != '簽': _line(doc, '受文者', c.get('recipient'))
    _line(doc, '發文日期', c['date'])
    if kind != '簽':
        _line(doc, '發文字號', c.get('doc_number'))
        _line(doc, '速別', c.get('speed'))
        doc.add_paragraph('密等及解密條件或保密期限：' + c.get('secrecy', ''))
        doc.add_paragraph('附件：' + ('、'.join(c['attachments']) if c.get('attachments') else ''))
    if kind == '開會通知單':
        m = c['meeting']
        _line(doc, '開會事由', c['subject']); _line(doc, '開會時間', m['time']); _line(doc, '開會地點', m['place'])
        _line(doc, '主持人', m['chair']); _line(doc, '聯絡人及電話', m['contact'])
        _line(doc, '出席者', '、'.join(m['attendees']))
        if m.get('observers'): _line(doc, '列席者', '、'.join(m['observers']))
        if m.get('agenda'): doc.add_paragraph('議程：'); _numbered(doc, m['agenda'])
        if c['explanation']: doc.add_paragraph('備註：'); _numbered(doc, c['explanation'])
    else:
        doc.add_paragraph('主旨：' + c['subject'])
        if kind == '公告':
            doc.add_paragraph('依據：'); _numbered(doc, c['basis'])
        elif c['explanation']:
            doc.add_paragraph('說明：'); _numbered(doc, c['explanation'])
        if c['proposal'] and c['proposal_label'] != '無':
            doc.add_paragraph(c['proposal_label'] + '：'); _numbered(doc, c['proposal'])
    if c.get('cc'):
        _line(doc, '正本', '、'.join(c['cc']['original']))
        doc.add_paragraph('副本：' + '、'.join(c['cc'].get('copy', [])))
    if kind == '簽' and c.get('sign_route'):
        doc.add_paragraph('陳核：' + '→'.join(c['sign_route']))
    doc.add_paragraph('')
    doc.add_heading('附註（不屬公文本文）', 1)
    doc.add_paragraph(f"行文方向：{c['direction']}　簽辦方式：{c['handling']}　模式：{c['mode']}")
    for source in d['sources']:
        doc.add_paragraph(f"依據來源：{source['title']}｜{source['status']}" + (f"｜{source['url']}" if source.get('url') else ''))
    if c.get('review_findings'):
        doc.add_heading('修改建議', 2)
        add_table(doc, ['位置', '問題', '原因', '建議寫法', '依據', '嚴重度'], [[f['location'], f['issue'], f['why'], f['fix'], f['basis'], f['severity']] for f in c['review_findings']], [1, 2, 3, 3, 2, 1])
    if c.get('pending'):
        doc.add_heading('待確認事項', 2)
        for x in c['pending']: doc.add_paragraph(x, style='List Bullet')


def render_conduct(doc, d):
    c = d['content']
    doc.add_paragraph(f"學制：{c['school_level']}　依行為事實記錄，不作綜合評價、不轉換等第。")
    for s in c['students']:
        doc.add_heading(f"學生代碼：{s['student_id']}", 1)
        add_table(doc, ['證據編號', '面向', '日期', '行為事實'], [[e['id'], e['aspect'], e['date'], e['observation']] for e in s['evidence']], [1, 1, 1, 5])
        doc.add_paragraph('評語：' + s['comment'])
        doc.add_paragraph('具體建議：' + s['suggestion'])
        doc.add_paragraph('引用證據：' + '、'.join(s['evidence_ids']))
        if s.get('parent_version'): doc.add_paragraph('給家長的版本：' + s['parent_version'])
